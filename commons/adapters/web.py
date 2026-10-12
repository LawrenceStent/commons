"""The web, for agents: the only way anything in a society reaches the internet (model calls aside).

This is the gate's second enforcement point, independent of the first (commons/application/gate.py, in the tool layer). Even if
the gate approved something by mistake, this layer refuses:

    - any host not on the egress allowlist (exact, or a subdomain of an entry written "*.example.org")
    - anything but https, IP-literal hosts, and hosts that resolve to private, loopback or link-local addresses
    - redirects to anywhere the allowlist doesn't cover (each hop is checked; at most 3)
    - pages a site's robots.txt disallows for us (search APIs, which are meant for programs, skip robots)
    - anything that isn't text (html, plain, json, xml), a PDF or a spreadsheet, and bodies over `max_bytes`
      (`pdf_max_bytes` for a PDF or spreadsheet). A PDF's text is extracted page by page, each page marked "[page N]";
      a spreadsheet (csv, ods, xlsx) becomes rows of cells joined by " | ", each sheet marked "[sheet NAME]"

Reads only: GET, no cookies, no credentials, a User-Agent that says what we are. Page text is untrusted: whoever
shows it to a model fences it as reference material.
"""

from __future__ import annotations

import html
import io
import ipaddress
import json
import re
import socket
import urllib.error
import urllib.parse
import urllib.request
import urllib.robotparser
from collections.abc import Callable
from html.parser import HTMLParser

from commons.adapters import secrets
from commons.application.ports import Page, SearchResult, WebError, host_of

USER_AGENT = "CommonsResearchBot/0.1 (a small research simulation; read-only, rate-limited)"
TEXT_TYPES = ("text/html", "text/plain", "application/json", "application/xml", "text/xml", "application/xhtml+xml")
PDF = "application/pdf"
SHEETS = {"text/csv": "csv", "application/csv": "csv", "application/vnd.oasis.opendocument.spreadsheet": "ods",
          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx"}
MAX_ROWS = 3000  # per sheet

# (url, headers) -> (status, headers, body). The default is urllib without automatic redirects; tests pass a fake.
Transport = Callable[[str, dict], tuple[int, dict, bytes]]


class EgressDenied(WebError):
    pass


# ── the allowlist ──────────────────────────────────────────────
def _private(host: str) -> bool:
    try:
        infos = socket.getaddrinfo(host, 443, proto=socket.IPPROTO_TCP)
    except socket.gaierror as e:
        raise WebError(f"can't resolve {host}") from e
    for *_, addr in infos:
        ip = ipaddress.ip_address(addr[0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast or ip.is_unspecified:
            return True
    return False


class Egress:
    def __init__(self, allow_hosts, resolve: Callable[[str], bool] | None = None):
        self.allow = frozenset(h.lower().strip().rstrip(".") for h in allow_hosts if h.strip())
        self._private = resolve or _private  # tests pass a stub: no DNS in the suite

    def allows_host(self, host: str) -> bool:
        return host in self.allow or any(a.startswith("*.") and (host == a[2:] or host.endswith(a[1:])) for a in self.allow)

    def check(self, url: str) -> str:
        """The url's host, or EgressDenied."""
        parts = urllib.parse.urlsplit(url)
        host = (parts.hostname or "").lower().rstrip(".")
        if parts.scheme != "https":
            raise EgressDenied(f"only https is allowed ({url[:80]})")
        if not host:
            raise EgressDenied(f"no host in {url[:80]}")
        if parts.username or parts.password:
            raise EgressDenied("urls with credentials are not allowed")
        try:
            ipaddress.ip_address(host)
            raise EgressDenied("addresses must be host names, not IPs")
        except ValueError:
            pass
        if not self.allows_host(host):
            raise EgressDenied(f"{host} is not on the operator's allowlist")
        if self._private(host):
            raise EgressDenied(f"{host} resolves to a private address")
        return host


# ── fetching ───────────────────────────────────────────────────
class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **kw):
        return None  # we follow redirects ourselves, checking each hop


def urllib_transport(timeout: float = 15.0, max_bytes: int = 2_000_000) -> Transport:
    opener = urllib.request.build_opener(_NoRedirect)

    def send(url: str, headers: dict) -> tuple[int, dict, bytes]:
        req = urllib.request.Request(url, headers=headers, method="GET")
        try:
            with opener.open(req, timeout=timeout) as r:
                return r.status, {k.lower(): v for k, v in r.headers.items()}, r.read(max_bytes + 1)
        except urllib.error.HTTPError as e:  # 3xx arrive here too, with no redirect handler
            return e.code, {k.lower(): v for k, v in (e.headers or {}).items()}, e.read(max_bytes + 1) if e.fp else b""
        except (urllib.error.URLError, TimeoutError, OSError) as e:
            raise WebError(f"couldn't reach {host_of(url)}: {getattr(e, 'reason', e)}") from e

    return send


class _Text(HTMLParser):
    SKIP = {"script", "style", "nav", "footer", "header", "noscript", "svg", "form", "aside"}
    BLOCK = {"p", "div", "li", "br", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article", "table"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.title, self._skip, self._in_title = [], "", 0, False
        self.links: list[tuple[str, str]] = []  # (href, its text), in page order
        self._href: str | None = None
        self._anchor: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag == "a" and not self._skip:
            self._href, self._anchor = dict(attrs).get("href"), []
        if tag in self.SKIP:
            self._skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag == "a" and self._href:
            self.links.append((self._href, " ".join("".join(self._anchor).split())))
            self._href = None
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag == "title":
            self._in_title = False
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self._skip:
            self.out.append(data)
            if self._href:
                self._anchor.append(data)


DOCUMENTS = (".pdf", ".csv", ".xlsx", ".xls", ".ods", ".json", ".zip")
MAX_LINKS = 40


def html_to_text(html: str, base: str = "") -> tuple[str, str]:
    """A page's title and text. With `base` (the page's url), the text ends with the page's links, so an agent can
    follow them: documents (PDFs, spreadsheets) first, then the rest, at most MAX_LINKS."""
    p = _Text()
    p.feed(html)
    text = re.sub(r"[ \t\r\f\v]+", " ", "".join(p.out))
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if base and (links := page_links(p.links, base)):
        text += "\n\nLinks on this page:\n" + "\n".join(f"- {label}: {url}" if label else f"- {url}" for url, label in links)
    return p.title.strip(), text


def page_links(raw: list[tuple[str, str]], base: str) -> list[tuple[str, str]]:
    seen, docs, rest = {base.split("#")[0]}, [], []
    for href, label in raw:
        url = urllib.parse.urljoin(base, href.strip()).split("#")[0]
        if not url.startswith(("https://", "http://")) or url in seen:
            continue
        seen.add(url)
        (docs if urllib.parse.urlsplit(url).path.lower().endswith(DOCUMENTS) else rest).append((url, label[:100]))
    return (docs + rest)[:MAX_LINKS]


_GENERIC = re.compile(r"^(powerpoint presentation|presentation|untitled|document|slide \d+|microsoft (word|powerpoint|"
                      r"excel)\b.*|.*\.(docx?|pptx?|xlsx?|indd|pdf))$", re.IGNORECASE)


def _generic(title: str) -> bool:
    """A PDF title that says nothing about the document (12 Oct: Ofwat's reports are all "PowerPoint Presentation")."""
    return not title or bool(_GENERIC.match(title.strip()))


def pdf_to_text(body: bytes, max_pages: int = 400) -> tuple[str, str]:
    """A PDF's title and text, each page headed "[page N]"; scanned pages (no text layer) come out empty."""
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        reader = PdfReader(io.BytesIO(body))
        title = str((reader.metadata or {}).get("/Title") or "").strip()
        pages = [f"[page {i + 1}]\n{(page.extract_text() or '').strip()}" for i, page in enumerate(reader.pages[:max_pages])]
        if _generic(title):  # "PowerPoint Presentation", "Microsoft Word - draft3.docx": the first real line instead
            first = next((line.strip() for page in pages for line in page.splitlines()[1:] if len(line.strip()) > 12), "")
            title = first[:120]
    except (PdfReadError, ValueError, KeyError, TypeError) as e:
        raise WebError(f"couldn't read the PDF: {e}") from e
    text = re.sub(r"[ \t\r\f\v]+", " ", "\n\n".join(pages))
    return title, re.sub(r"\n\s*\n+", "\n\n", text).strip()


def _sheet_kind(ctype: str, url: str) -> str | None:
    if ctype in SHEETS:
        return SHEETS[ctype]
    ext = urllib.parse.urlsplit(url).path.lower().rsplit(".", 1)[-1]
    return ext if ext in ("csv", "ods", "xlsx") and ctype in ("application/octet-stream", "binary/octet-stream",
                                                                 "application/zip", "text/plain", "") else None


def _rows_text(name: str, rows) -> str:
    lines = []
    for row in rows:
        cells = ["" if c is None else " ".join(str(c).split()) for c in row]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            lines.append(" | ".join(cells))
        if len(lines) >= MAX_ROWS:
            lines.append(f"(more rows not shown: the first {MAX_ROWS} only)")
            break
    return f"[sheet {name}]\n" + "\n".join(lines) if lines else ""


def _ods_rows(body: bytes):
    """The sheets of an OpenDocument spreadsheet: a zip whose content.xml holds table:table elements."""
    import zipfile
    from xml.etree import ElementTree as ET
    ns = {"table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0", "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0"}
    root = ET.fromstring(zipfile.ZipFile(io.BytesIO(body)).read("content.xml"))
    rep = f"{{{ns['table']}}}number-columns-repeated"
    for table in root.iter(f"{{{ns['table']}}}table"):
        rows = []
        for tr in table.iter(f"{{{ns['table']}}}table-row"):
            row = []
            for cell in tr:
                value = " ".join("".join(p.itertext()) for p in cell.iter(f"{{{ns['text']}}}p"))
                row += [value] * min(int(cell.get(rep, "1")), 50 if value else 1)  # repeated blanks pad wide rows
            rows.append(row)
            if len(rows) > MAX_ROWS:
                break
        yield table.get(f"{{{ns['table']}}}name", "sheet"), rows


def sheet_to_text(body: bytes, kind: str) -> str:
    """A spreadsheet as text: each sheet's non-empty rows, cells joined by " | "."""
    import csv
    import zipfile
    try:
        if kind == "csv":
            return _rows_text("csv", csv.reader(io.StringIO(body.decode("utf-8-sig", "replace"))))
        if kind == "xlsx":
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            return "\n\n".join(t for ws in book.worksheets if (t := _rows_text(ws.title, ws.iter_rows(values_only=True))))
        return "\n\n".join(t for name, rows in _ods_rows(body) if (t := _rows_text(name, rows)))
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
        raise WebError(f"couldn't read the spreadsheet: {e}") from e


class Fetcher:
    def __init__(self, egress: Egress, transport: Transport | None = None, max_bytes: int = 2_000_000,
                 max_chars: int = 20_000, pdf_max_bytes: int = 30_000_000, pdf_max_chars: int = 200_000):
        self.egress, self.max_bytes, self.max_chars = egress, max_bytes, max_chars
        self.pdf_max_bytes, self.pdf_max_chars = pdf_max_bytes, pdf_max_chars
        self.transport = transport or urllib_transport(max_bytes=max(max_bytes, pdf_max_bytes))
        self._robots: dict[str, urllib.robotparser.RobotFileParser] = {}
        self.requests = 0  # every request that reached the transport, robots.txt included

    def _get(self, url: str) -> tuple[str, dict, bytes]:
        for _ in range(4):
            self.egress.check(url)  # every hop, before anything is sent
            self.requests += 1
            status, headers, body = self.transport(url, {"User-Agent": USER_AGENT,
                                                        "Accept": "text/html,text/plain,application/json,application/pdf,"
                                                                  "text/csv,application/vnd.oasis.opendocument.spreadsheet,"
                                                                  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
            if status in (301, 302, 303, 307, 308) and headers.get("location"):
                url = urllib.parse.urljoin(url, headers["location"])
                continue
            if status >= 400:
                raise WebError(f"{host_of(url)} answered {status}")
            ctype = headers.get("content-type", "").split(";")[0].strip().lower()
            limit = self.pdf_max_bytes if ctype == PDF or _sheet_kind(ctype, url) else self.max_bytes
            if len(body) > limit:
                raise WebError(f"the page at {url[:80]} is larger than {limit} bytes")
            return url, headers, body
        raise WebError("too many redirects")

    def _allowed_by_robots(self, url: str) -> bool:
        parts = urllib.parse.urlsplit(url)
        root = f"{parts.scheme}://{parts.netloc}"
        rp = self._robots.get(root)
        if rp is None:
            rp = urllib.robotparser.RobotFileParser()
            try:
                _, _, body = self._get(f"{root}/robots.txt")
                rp.parse(body.decode("utf-8", "replace").splitlines())
            except WebError:
                rp.parse([])  # no robots.txt (or unreadable): the site hasn't said no
            self._robots[root] = rp
        return rp.can_fetch(USER_AGENT, url)

    def fetch(self, url: str, robots: bool = True) -> Page:
        self.egress.check(url)
        if robots and not self._allowed_by_robots(url):
            raise WebError(f"{host_of(url)}'s robots.txt asks bots not to read {urllib.parse.urlsplit(url).path}")
        final, headers, body = self._get(url)
        ctype = headers.get("content-type", "text/html").split(";")[0].strip().lower()
        if ctype == PDF or body[:5] == b"%PDF-":
            title, text = pdf_to_text(body)
            if not text.replace("[page", "").strip(" \n]0123456789"):
                raise WebError(f"{url[:80]} is a PDF with no text layer (a scan)")
            return Page(final, title or final.rsplit("/", 1)[-1], text[: self.pdf_max_chars])
        if kind := _sheet_kind(ctype, final):
            text = sheet_to_text(body, kind)
            if not text:
                raise WebError(f"{url[:80]} is a spreadsheet with nothing in it")
            return Page(final, final.rsplit("/", 1)[-1], text[: self.pdf_max_chars])
        if ctype not in TEXT_TYPES:
            raise WebError(f"{url[:80]} is {ctype}, not text, a PDF or a spreadsheet")
        raw = body.decode("utf-8", "replace")
        if "html" not in ctype:
            return Page(final, final, raw[: self.max_chars])
        title, text = html_to_text(raw, base=final)
        body_text, _, links = text.partition("\n\nLinks on this page:\n")
        links = f"\n\nLinks on this page:\n{links}" if links else ""  # kept whole: the page is cut, not its links
        return Page(final, title or final, body_text[: max(0, self.max_chars - len(links))] + links)

    def text(self, url: str) -> str:
        """An API call's body as text (a price feed's CSV, say). APIs are meant for programs: robots.txt doesn't apply."""
        _, _, body = self._get(url)
        return body.decode("utf-8", "replace")

    def json(self, url: str) -> dict:
        """A JSON API call (search providers). APIs are meant for programs, so robots.txt doesn't apply."""
        _, _, body = self._get(url)
        try:
            return json.loads(body)
        except json.JSONDecodeError as e:
            raise WebError(f"{host_of(url)} didn't return JSON") from e


# ── search ─────────────────────────────────────────────────────
class WikipediaSearch:
    """Search through Wikipedia's public API: free, no key, and its host must still be on the allowlist."""

    def __init__(self, fetcher: Fetcher, lang: str = "en"):
        self.fetcher, self.host = fetcher, f"{lang}.wikipedia.org"

    def search(self, query: str, k: int = 5) -> list[SearchResult]:
        q = urllib.parse.urlencode({"action": "query", "list": "search", "srsearch": query[:200], "srlimit": k,
                                    "format": "json", "utf8": 1})
        data = self.fetcher.json(f"https://{self.host}/w/api.php?{q}")
        out = []
        for hit in data.get("query", {}).get("search", [])[:k]:
            title = hit.get("title", "")
            snippet = html.unescape(re.sub(r"<[^>]+>", "", hit.get("snippet", "")))
            out.append(SearchResult(title, f"https://{self.host}/wiki/{urllib.parse.quote(title.replace(' ', '_'))}", snippet))
        return out


class GovUkSearch:
    """GOV.UK's public search API: free, no key; everything published on gov.uk (departments, the Environment Agency,
    Ofwat's publications, statistics). Its host must be on the allowlist like any other."""
    host = "www.gov.uk"

    def __init__(self, fetcher: Fetcher):
        self.fetcher = fetcher

    def search(self, query: str, k: int = 5) -> list[SearchResult]:
        q = urllib.parse.urlencode([("q", query[:200]), ("count", k), ("fields", "title"), ("fields", "link"),
                                    ("fields", "description")])
        data = self.fetcher.json(f"https://{self.host}/api/search.json?{q}")
        out = []
        for hit in data.get("results", [])[:k]:
            link = str(hit.get("link", ""))
            url = link if link.startswith("https://") else f"https://{self.host}{link}"
            out.append(SearchResult(str(hit.get("title", "")), url, str(hit.get("description", ""))))
        return out


# (url, headers, body) -> (status, body): Tavily is a POST with a key, which the read-only Fetcher never sends
Poster = Callable[[str, dict, bytes], tuple[int, bytes]]


def urllib_post(url: str, headers: dict, body: bytes) -> tuple[int, bytes]:
    req = urllib.request.Request(url, data=body, headers=headers, method="POST")
    try:
        with urllib.request.build_opener(_NoRedirect).open(req, timeout=20) as r:
            return r.status, r.read(2_000_000)
    except urllib.error.HTTPError as e:
        return e.code, e.read()[:2000] if e.fp else b""
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        raise WebError(f"couldn't reach {host_of(url)}: {getattr(e, 'reason', e)}") from e


class TavilySearch:
    """Tavily's web search API (TAVILY_API_KEY; 1,000 free searches a month). The whole web, so its results can name
    pages on hosts the allowlist doesn't cover: those can't be read until you add the host."""
    host = "api.tavily.com"

    def __init__(self, fetcher: Fetcher, key: str | None = None, post: Poster = urllib_post):
        self.fetcher, self.key, self.post = fetcher, key, post

    def search(self, query: str, k: int = 5) -> list[SearchResult]:
        key = self.key or secrets.get("TAVILY_API_KEY")
        if not key:
            raise WebError("no TAVILY_API_KEY set")
        url = f"https://{self.host}/search"
        self.fetcher.egress.check(url)  # the same allowlist as every read
        self.fetcher.requests += 1
        body = json.dumps({"query": query[:400], "max_results": k, "search_depth": "basic"}).encode()
        status, data = self.post(url, {"Content-Type": "application/json", "Authorization": f"Bearer {key}",
                                       "User-Agent": USER_AGENT}, body)
        if status >= 400:
            raise WebError(f"Tavily answered {status}")
        try:
            results = json.loads(data).get("results", [])
        except json.JSONDecodeError as e:
            raise WebError("Tavily didn't return JSON") from e
        return [SearchResult(str(r.get("title", "")), str(r.get("url", "")), str(r.get("content", ""))) for r in results[:k]]


class MultiSearch:
    """Several providers asked in turn (`[gate] search = "govuk+tavily"`), their results interleaved without repeats.
    A provider that fails is skipped, as long as another answers."""

    def __init__(self, providers: list):
        self.providers = providers

    @property
    def host(self) -> str:
        return self.providers[0].host

    def search(self, query: str, k: int = 5) -> list[SearchResult]:
        lists, errors = [], []
        for p in self.providers:
            try:
                lists.append(p.search(query, k))
            except WebError as e:
                errors.append(f"{p.host}: {e}")
        if not lists:
            raise WebError("; ".join(errors) or "no search provider")
        out, seen = [], set()
        for rank in range(k):
            for results in lists:
                if rank < len(results) and results[rank].url not in seen:
                    seen.add(results[rank].url)
                    out.append(results[rank])
        return out[: k * len(lists)]


SEARCHERS = {"wikipedia": WikipediaSearch, "govuk": GovUkSearch, "tavily": TavilySearch}


def searcher(spec: str, fetcher: Fetcher):
    """The search a `[gate] search` setting names: one provider, several joined with "+", or none."""
    names = [n for n in spec.split("+") if n and n != "none"]
    providers = [SEARCHERS[n](fetcher) for n in names]
    return None if not providers else providers[0] if len(providers) == 1 else MultiSearch(providers)


class WebAccess:
    """The web port (commons.application.ports.WebPort): a fetcher behind the allowlist, and an optional search provider. The
    allowlist follows the operator's [gate] policy, re-read every cycle."""

    def __init__(self, fetcher: Fetcher, searcher=None):
        self.fetcher, self.searcher = fetcher, searcher

    @property
    def search_host(self) -> str | None:
        return getattr(self.searcher, "host", None) if self.searcher is not None else None

    def set_hosts(self, hosts) -> None:
        self.fetcher.egress.allow = frozenset(h.lower().strip().rstrip(".") for h in hosts if h.strip())

    def fetch(self, url: str) -> Page:
        return self.fetcher.fetch(url)

    def search(self, query: str, k: int = 5) -> list[SearchResult]:
        if self.searcher is None:
            raise WebError("this society has no web search")
        return self.searcher.search(query, k)

    @classmethod
    def default(cls, search: str = "wikipedia") -> WebAccess:
        fetcher = Fetcher(Egress(()))  # no hosts until the operator allows some
        return cls(fetcher, searcher(search, fetcher))
